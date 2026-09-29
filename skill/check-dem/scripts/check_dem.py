"""Check DEM: คำนวณวันเริ่มเก็บค่าใช้จ่าย (DEM) และวัน REC. จาก Freedays + ATA

กติกา
  DEM  = ATA + (เลขตัวหน้าของ Freedays + เลขตัวหน้าหลัง ADD)
         เช่น "7/5 Add 7/9" -> 7 + 7 = 14 วัน ; ATA 27/09/2026 -> 2026-10-11
  TOTAL = จำนวนวันที่ได้จาก Freedays
  REC. = ATA + 3 วัน  (เฉพาะแถวที่มีคำว่า TOS) ; ถ้าไม่มี TOS เว้นว่าง
  DEM (และช่อง Freedays) วันเดียวกันได้สีเดียวกัน ไล่สีจากวันแรกสุด (แดง/ส้ม) ไปวันหลังสุด (เขียว/ฟ้า)

ผลลัพธ์บันทึกเป็น output\\Check DEM.xlsx

วิธีใช้
  python check_dem.py <ไฟล์ .xls/.xlsx> <ATA dd/mm/yyyy> [โฟลเดอร์ผลลัพธ์]
  python check_dem.py input\\DEM-KMGY.xls 27/09/2026

ตัวเลือก
  --cols C-F       ใช้เฉพาะคอลัมน์ C ถึง F ของต้นฉบับ (แบบฟอร์ม SUR: B/L No. | Freedays | RECEIPT B/L | Consignee)
  --name START_DEM ตั้งชื่อไฟล์ผลลัพธ์ (ได้ START_DEM.xlsx) แทน "Check DEM.xlsx"
  python check_dem.py input\\SUR.xls 27/09/2026 input --cols C-F --name START_DEM
"""
import colorsys
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

TOS_EXTRA_DAYS = 3
OUTPUT_NAME = "Check DEM.xlsx"

FONT = "Tahoma"
NAVY = "1F3864"
LINE = "BFBFBF"


def free_days(text):
    """'7/5 Add 7/9' -> 14 ; '3/3' -> 3 ; ไม่พบตัวเลข -> None"""
    text = str(text or "")
    parts = re.split(r"\badd\b", text, flags=re.I)
    total, found = 0, False
    for part in parts[:2]:
        m = re.search(r"\d+", part)
        if m:
            total += int(m.group())
            found = True
    return total if found else None


def find_col(header, name):
    for i, h in enumerate(header):
        if name.lower() in str(h).lower():
            return i
    return None


def date_palette(dates):
    """วันที่เรียงจากเร็วสุด -> ช้าสุด ได้สีไล่จากแดงอ่อน -> ส้ม -> เหลือง -> เขียว -> ฟ้า"""
    dates = sorted(set(dates))
    n = len(dates)
    colors = {}
    for i, d in enumerate(dates):
        hue = 0.0 if n == 1 else (i / (n - 1)) * 0.55  # 0 = แดง ... 0.55 = ฟ้า
        r, g, b = colorsys.hls_to_rgb(hue, 0.80, 0.75)
        colors[d] = f"{int(r * 255):02X}{int(g * 255):02X}{int(b * 255):02X}"
    return colors


def pop_option(args, name):
    """ดึงค่า --name value ออกจาก args ; ไม่มี -> None"""
    if name in args:
        i = args.index(name)
        value = args[i + 1]
        del args[i:i + 2]
        return value
    return None


def main():
    cols = pop_option(sys.argv, "--cols")
    out_name = pop_option(sys.argv, "--name") or OUTPUT_NAME
    if not out_name.lower().endswith(".xlsx"):
        out_name += ".xlsx"
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = Path(sys.argv[1])
    ata_text = sys.argv[2] if len(sys.argv) > 2 else input("ATA (dd/mm/yyyy): ").strip()
    ata = datetime.strptime(ata_text, "%d/%m/%Y")

    raw = pd.read_excel(src, header=None, dtype=str).fillna("")
    if cols:
        first, last = (column_index_from_string(c.strip()) - 1 for c in re.split(r"[-:]", cols.upper()))
        raw = raw.iloc[:, first:last + 1]
    header = [str(h).strip() for h in raw.iloc[0]]
    rows = raw.iloc[1:].values.tolist()

    c_free = find_col(header, "Freedays")
    if c_free is None:
        sys.exit("ไม่พบคอลัมน์ Freedays ในแถวแรกของไฟล์")
    others = [i for i in range(len(header)) if i != c_free]

    # ---- คำนวณ ----
    data = []
    for r in rows:
        if not any(str(v).strip() for v in r):
            continue
        fd_text = str(r[c_free]).strip()
        days = free_days(fd_text)
        dem = ata + timedelta(days=days) if days is not None else None
        has_tos = any(re.search(r"\bTOS\b", str(v), re.I) for v in r)
        rec = ata + timedelta(days=TOS_EXTRA_DAYS) if has_tos else None
        data.append((fd_text, days, dem, rec, [str(r[i]).strip() for i in others]))

    colors = date_palette(d[2] for d in data if d[2] is not None)

    # ---- สร้างไฟล์ ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Check DEM"
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color=LINE)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", indent=1)

    out_header = ["No.", "Freedays", "TOTAL", "DEM.", "REC."] + [header[i] for i in others]
    ncol = len(out_header)
    last_col = get_column_letter(ncol)

    # ชื่อรายงาน
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = "CHECK DEM"
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30

    # ข้อมูลไฟล์ (แถวเดียวใต้ชื่อรายงาน)
    info = [("ต้นฉบับ", src.name, None), ("ATA", ata, "DD/MM/YYYY"), ("จำนวน B/L", len(data), None)]
    for i, (label, value, fmt) in enumerate(info):
        a = ws.cell(row=2, column=1 + i * 2, value=label)
        b = ws.cell(row=2, column=2 + i * 2, value=value)
        a.font = Font(name=FONT, bold=True, color=NAVY)
        a.fill = PatternFill("solid", fgColor="D9E1F2")
        a.alignment = left
        b.font = Font(name=FONT, bold=(label == "ATA"), size=12 if label == "ATA" else 10)
        b.alignment = center if label == "ATA" else left
        for c in (a, b):
            c.border = box
        if fmt:
            b.number_format = fmt
    ws.row_dimensions[2].height = 22

    # หัวตาราง
    hr = 4
    ws.row_dimensions[hr].height = 30
    for col, h in enumerate(out_header, 1):
        c = ws.cell(row=hr, column=col, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
    receipt_cols = {i for i, h in enumerate(out_header, 1) if "RECEIPT" in h.upper()}

    # รายการ
    problems = 0
    for n, (fd_text, days, dem, rec, rest) in enumerate(data, 1):
        row = hr + n
        ws.row_dimensions[row].height = 20
        values = [n, fd_text, days, dem if dem else "CHECK", rec] + rest
        stripe = PatternFill("solid", fgColor="F7F9FC" if n % 2 == 0 else "FFFFFF")
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v if v != "" else None)
            c.font = Font(name=FONT)
            c.fill, c.border = stripe, box
            h = out_header[col - 1]
            c.alignment = left if h in ("Consignee", "Freedays") else center

        ws.cell(row=row, column=3).font = Font(name=FONT, bold=True)
        dem_cell = ws.cell(row=row, column=4)
        dem_cell.number_format = "YYYY-MM-DD"
        dem_cell.font = Font(name=FONT, bold=True)
        if dem:
            dem_cell.fill = PatternFill("solid", fgColor=colors[dem])
            ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=colors[dem])
        else:
            dem_cell.fill = PatternFill("solid", fgColor="FF0000")
            dem_cell.font = Font(name=FONT, bold=True, color="FFFFFF")
            problems += 1

        rec_cell = ws.cell(row=row, column=5)
        rec_cell.number_format = "DD/MM/YYYY"
        if rec:
            rec_cell.font = Font(name=FONT, bold=True, color="C00000")
            rec_cell.fill = PatternFill("solid", fgColor="FCE4D6")
            ws.cell(row=row, column=2).font = Font(name=FONT, bold=True, color="C00000")

        for col in receipt_cols:
            c = ws.cell(row=row, column=col)
            c.font = Font(name=FONT, bold=True, color="C00000")
            c.fill = PatternFill("solid", fgColor="FFF2CC")

    # ขนาดคอลัมน์ให้พอดีกับตัวอักษรที่ยาวที่สุดในแต่ละช่อง
    for col in range(1, ncol + 1):
        longest = 0
        # หัวตารางตัดบรรทัดได้ จึงนับแค่คำที่ยาวที่สุดของหัวตาราง
        longest = max(len(w) for w in out_header[col - 1].split())
        # นับทั้งกล่องข้อมูล ด้านบนและรายการ ; ตัวหนา/ตัวใหญ่ต้องใช้ที่มากขึ้น
        for row in range(2, hr + len(data) + 1):
            if row == hr:
                continue
            c = ws.cell(row=row, column=col)
            if c.value is None:
                continue
            text = c.value.strftime("%d/%m/%Y") if hasattr(c.value, "strftime") else str(c.value)
            scale = (c.font.sz or 10) / 10 * (1.12 if c.font.b else 1.0)
            longest = max(longest, len(text) * scale)
        ws.column_dimensions[get_column_letter(col)].width = longest * 1.15 + 3
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{last_col}{hr + len(data)}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"

    out_dir = Path(sys.argv[3]) if len(sys.argv) > 3 else src.resolve().parent.parent / "output"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / out_name
    wb.save(out)
    print(f"บันทึกแล้ว: {out}  ({len(data)} แถว, อ่าน Freedays ไม่ได้ {problems} แถว)")


if __name__ == "__main__":
    main()
